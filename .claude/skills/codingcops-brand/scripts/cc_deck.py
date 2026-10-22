#!/usr/bin/env python3
"""
cc_deck.py — build on-brand CodingCops PowerPoint decks.

Implements the slide types in guidelines/08-presentations.md on a 1920 x 1080
design canvas (13.333 x 7.5 in PPTX; 1 px = 0.5 pt = 6350 EMU), using the exact
values in tokens/brand-tokens.json and the files in assets/.

Use from Python:

    import sys; sys.path.insert(0, "<skill>/scripts")
    from cc_deck import Deck
    d = Deck()
    d.title(eyebrow="Pitch deck", headline="Software built with", headline_2="brilliance and precision",
            script="for Acme", period="Q4 '26")
    d.section(1, "The problem", script="why now")
    d.big_number("40%", "Releases slip by 40% when QA starts late", context="Industry survey, 2025")
    d.bullets("Approach", "We embed QA from the first sprint",
              ["Test plans written with the user stories", "Automated checks on every merge",
               "Weekly quality report for stakeholders"])
    d.kpi_chart("Results", "Defects fell 62% in six months",
                kpis=[{"value": "62%", "label": "Fewer defects", "context": "▼ down vs. Q1", "highlight": True},
                      {"value": "3.1x", "label": "Faster releases", "context": "▲ up from monthly"}],
                chart={"categories": ["Jan", "Feb", "Mar", "Apr", "May", "Jun"],
                       "values": [120, 104, 88, 70, 58, 46], "highlight": 5},
                note="June: lowest defect count on record")
    d.cards("Services", "Four ways we help you ship", cards=[{"title": "...", "body": "..."}, ...])
    d.closing(script="let's build", contact="[contact email] · codingcops.com")
    d.save("acme-pitch.pptx")

Or from a JSON spec (list of slides, each {"type": ..., **arguments}):

    python3 cc_deck.py --spec deck.json --out deck.pptx [--preview previews/]
    python3 cc_deck.py --demo demo.pptx --preview previews/
"""
from __future__ import annotations

import argparse
import copy
import json
import math
import shutil
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

SKILL = Path(__file__).resolve().parents[1]
TOKENS = json.loads((SKILL / "tokens" / "brand-tokens.json").read_text(encoding="utf-8"))
LOGO_DIR = SKILL / TOKENS["logos"]["$dir"]
BLOB = json.loads((SKILL / "assets" / "shapes" / "emerald-blob.json").read_text(encoding="utf-8"))

W, H = 1920, 1080          # design canvas, px
M = 96                     # outer margin, px
CONTENT_TOP = 380          # content starts under eyebrow + headline
FOOTER_Y = 984             # footer row top (64 px breathing room above it)
CONTENT_BOTTOM = FOOTER_Y - 64

DARK = {"prussian-blue", "full-spectrum-blue"}
SECTION_TONES = ["full-spectrum-blue", "emerald", "prussian-blue"]
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def E(px: float) -> Emu:
    """Canvas px -> EMU."""
    return Emu(int(round(px * 6350)))


def hexof(token: str) -> str:
    return TOKENS["colors"][token]["hex"].lstrip("#").upper()


def rgb(token: str) -> RGBColor:
    return RGBColor.from_string(hexof(token))


def text_on(fill: str) -> str:
    """Default text colour for a fill, per text_on_fill rules."""
    return TOKENS["text_on_fill"].get(fill, {}).get("allowed", ["prussian-blue"])[0]


def font_for(weight: int, family: str = "sans") -> tuple[str, bool]:
    if family == "accent":
        return TOKENS["typography"]["families"]["accent"]["family"], False
    names = TOKENS["typography"]["families"]["sans"]["office_names"]
    bold_for = TOKENS["typography"]["families"]["sans"]["office_bold_flag_for"]
    return names[str(weight)], weight in bold_for


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = E(W), E(H)
        self._brand_theme()
        self._layout = self.prs.slide_layouts[6]  # blank
        self._section = ""
        self._section_count = 0
        self._missing: set[str] = set()

    # ------------------------------------------------------------------ setup
    def _brand_theme(self):
        """Brand colours in PowerPoint's colour picker, Satoshi as theme font,
        and 4:3 master placeholders stretched to 16:9."""
        master = self.prs.slide_master
        theme = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        ns = {"a": A_NS}
        scheme = {
            "dk1": "prussian-blue", "lt1": "white", "dk2": "prussian-blue", "lt2": "platinum",
            "accent1": "emerald", "accent2": "full-spectrum-blue", "accent3": "vibrant-coral",
            "accent4": "slate", "accent5": "emerald-wash", "accent6": "blue-wash",
            "hlink": "full-spectrum-blue", "folHlink": "full-spectrum-blue",
        }
        clr = root.find(".//a:clrScheme", ns)
        clr.set("name", "CodingCops")
        for tag, token in scheme.items():
            el = clr.find(f"a:{tag}", ns)
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{{{A_NS}}}srgbClr", val=hexof(token))
        fs = root.find(".//a:fontScheme", ns)
        fs.set("name", "CodingCops")
        fs.find("a:majorFont/a:latin", ns).set("typeface", "Satoshi")
        fs.find("a:minorFont/a:latin", ns).set("typeface", "Satoshi")
        theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        k = 1920 / 1440  # default template is 10in wide; slides are 13.333in
        for container in [master, *master.slide_layouts]:
            for shp in container.shapes:
                if shp.left is not None:
                    shp.left, shp.width = Emu(int(shp.left * k)), Emu(int(shp.width * k))

    def _slide(self, ground: str, notes: str | None = None):
        s = self.prs.slides.add_slide(self._layout)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(ground)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # -------------------------------------------------------------- primitives
    def _text(self, s, x, y, w, h, text, *, size, weight=400, color="prussian-blue", family="sans",
              align="left", anchor="top", caps=False, tracking=0.0, line=1.1, rotation=0.0,
              outline=None, outline_px=2.0):
        tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        name, bold = font_for(weight, family)
        lines = text if isinstance(text, list) else [text]
        for i, ln in enumerate(lines):
            if isinstance(ln, str):
                ln = {"text": ln}
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = line
            r = p.add_run()
            r.text = ln["text"]
            f = r.font
            f.name, f.bold, f.size = name, bold, Pt(size * 0.5)
            rPr = r._r.get_or_add_rPr()
            if caps:
                rPr.set("cap", "all")
            if tracking:
                rPr.set("spc", str(int(round(tracking * 50))))
            if outline:
                ln_el = etree.Element(qn("a:ln"), w=str(int(outline_px * 6350)))
                sf = etree.SubElement(ln_el, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr"), val=hexof(outline))
                rPr.insert(0, ln_el)
                ln_el.addnext(etree.Element(qn("a:noFill")))
            else:
                f.color.rgb = rgb(ln.get("color", color))
        if rotation:
            tb.rotation = rotation
        return tb

    def _box(self, s, x, y, w, h, fill=None, outline="prussian-blue", stroke=2.0, radius=32.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        shp = s.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
        _strip_style(shp)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if outline:
            shp.line.color.rgb = rgb(outline)
            shp.line.width = E(stroke)
        else:
            shp.line.fill.background()
        return shp

    def _blob(self, s, x, y, h):
        """The one emerald blob (traced from the monogram's green half)."""
        w = h * BLOB["aspect_w_over_h"]
        pts = [(int(E(x + px * w)), int(E(y + py * h))) for px, py in BLOB["points_normalized"]]
        fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
        fb.add_line_segments(pts[1:], close=True)
        shp = fb.convert_to_shape()
        _strip_style(shp)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb("emerald")
        shp.line.fill.background()
        shp.name = "Emerald blob"
        return shp

    def _logo(self, s, key, x, y, *, w=None, h=None):
        info = TOKENS["logos"][key]
        path = LOGO_DIR / info["file"]
        if path.exists():
            pic = s.shapes.add_picture(str(path), E(x), E(y), width=E(w) if w else None, height=E(h) if h else None)
            pic.name = f"Logo {key}"
            return pic
        self._missing.add(info["file"])
        bw, bh = (w or 400), (h or 64)
        b = self._box(s, x, y, bw, bh, fill=None, outline="slate", stroke=2, radius=8)
        self._text(s, x, y, bw, bh, f"Logo: add {info['file']}", size=18, color="slate", align="center",
                   anchor="middle")
        return b

    def _eyebrow(self, s, text, dark=False, y=M, ground=None):
        """Eyebrow: full-spectrum-blue on light, emerald on prussian, white on
        full-spectrum-blue, prussian on emerald (text_on_fill rules)."""
        color = {"prussian-blue": "emerald", "full-spectrum-blue": "white",
                 "emerald": "prussian-blue"}.get(ground, "emerald" if dark else "full-spectrum-blue")
        self._text(s, M, y, 1200, 30, text, size=20, weight=500, caps=True, tracking=1, color=color)

    def _headline(self, s, text, *, x=M, y=132, w=W - 2 * M, size=72, color="prussian-blue"):
        self._text(s, x, y, w, size * 2.3, text, size=size, weight=700, tracking=-size * 0.015,
                   line=1.05, color=color)

    def _footer(self, s):
        page = len(self.prs.slides)
        mh = 32
        self._logo(s, "monogram", M, FOOTER_Y, h=mh)
        mono_w = mh * 875 / 540
        if self._section:
            self._text(s, M + mono_w + 20, FOOTER_Y + 4, 900, 28, self._section, size=18, weight=500,
                       caps=True, tracking=1, color="slate")
        self._text(s, W - M - 200, FOOTER_Y + 4, 200, 28, str(page), size=18, weight=500,
                   tracking=1, color="slate", align="right")

    def _outlined(self, s, text, right, top, *, size, color="white", stroke=3, tracking=-0.02):
        """Outlined (hollow) Satoshi Black text, e.g. the period marker and section
        numbers. Rendered as a transparent PNG so it looks identical in PowerPoint,
        Keynote, Google Slides and PDF exports (text outlines are not portable)."""
        from PIL import Image, ImageChops, ImageDraw, ImageFont
        import tempfile
        k = 2  # render at 2x for crisp edges
        font = ImageFont.truetype(str(SKILL / "assets" / "fonts" / "Satoshi-Black.otf"), int(size * k))
        sw = max(1, int(round(stroke * k)))
        adv = [font.getlength(ch) + tracking * size * k for ch in text]
        wpx, hpx = int(sum(adv) + 4 * sw + size * k * 0.2), int(size * k * 1.4 + 4 * sw)
        a, b = Image.new("L", (wpx, hpx), 0), Image.new("L", (wpx, hpx), 0)
        da, db = ImageDraw.Draw(a), ImageDraw.Draw(b)
        x = 2 * sw
        for ch, step in zip(text, adv):
            da.text((x, 2 * sw), ch, font=font, fill=255, stroke_width=sw, stroke_fill=255)
            db.text((x, 2 * sw), ch, font=font, fill=255)
            x += step
        mask = ImageChops.subtract(a, b)
        bbox = mask.getbbox()
        mask = mask.crop(bbox)
        rgba = Image.new("RGBA", mask.size, "#" + hexof(color))
        rgba.putalpha(mask)
        path = Path(tempfile.mkdtemp()) / "outlined.png"
        rgba.save(path)
        w_px, h_px = mask.size[0] / k, mask.size[1] / k
        pic = s.shapes.add_picture(str(path), E(right - w_px), E(top), E(w_px), E(h_px))
        pic.name = f"Outlined {text}"
        pic._element.nvPicPr.cNvPr.set("descr", text)
        return pic

    def _dash_list(self, s, items, x, y, w, *, size=32, gap=40, color="prussian-blue"):
        items = items[:3]
        for it in items:
            lines = max(1, math.ceil(len(it) * size * 0.5 / w))
            hgt = lines * size * 1.3
            self._box(s, x, y + size * 0.62, 28, 6, fill="emerald", outline=None, radius=3)
            self._text(s, x + 52, y, w - 52, hgt, it, size=size, color=color, line=1.25)
            y += hgt + gap
        return y

    # ------------------------------------------------------------ slide types
    def title(self, headline, headline_2=None, *, eyebrow=None, script=None, period=None,
              second_line_color="emerald", notes=None):
        """Opening slide: prussian ground, inline white logo, two-line headline,
        one script line, emerald blob, outlined period marker."""
        s = self._slide("prussian-blue", notes)
        self._logo(s, "inline-white", M, M, w=400)
        if period:
            self._outlined(s, period, W - M, 200, size=200, color="white", stroke=3)
        self._blob(s, W - 430, H - 470, 660)
        y = 560 if script else 620
        if eyebrow:
            self._eyebrow(s, eyebrow, y=y - 52, ground="prussian-blue")
        lines = [{"text": headline, "color": "white"}]
        if headline_2:
            lines.append({"text": headline_2, "color": second_line_color})
        self._text(s, M, y, 1240, 200, lines, size=88, weight=700, tracking=-1.3, line=1.04)
        if script:
            self._text(s, M + 8, y + 210, 900, 110, script, size=64, family="accent", color="white",
                       rotation=-3)
        return s

    def section(self, number, title, *, script=None, tone=None, notes=None):
        """Divider: rotating full-bleed colour, giant outlined number, title, one script line."""
        if tone is None:
            tone = SECTION_TONES[self._section_count % 3]
        self._section_count += 1
        self._section = title
        s = self._slide(tone, notes)
        ink = "white" if tone in DARK else "prussian-blue"
        if tone in DARK:
            self._blob(s, W - 400, H - 420, 600)
        self._outlined(s, f"{int(number):02d}", W - M, 110, size=400, color=ink, stroke=4)
        self._eyebrow(s, f"Section {int(number):02d}", y=470, ground=tone)
        self._text(s, M, 512, 1300, 300, title, size=140, weight=900, tracking=-3, line=1.0, color=ink)
        if script:
            lines = 2 if len(title) * 140 * 0.58 > 1300 else 1
            self._text(s, M + 8, 512 + lines * 140 + 20, 900, 110, script, size=64, family="accent",
                       color=ink, rotation=-3)
        return s

    def big_number(self, number, headline, *, context=None, eyebrow=None, notes=None):
        """One hero number and its takeaway. Use for the deck's 3–5 key facts."""
        s = self._slide("platinum", notes)
        if eyebrow:
            self._eyebrow(s, eyebrow)
        self._text(s, M, 300, 900, 300, number, size=260, weight=900, tracking=-6, line=1.0,
                   anchor="bottom")
        self._box(s, M + 8, 628, 200, 16, fill="emerald", outline=None, radius=8)
        self._text(s, 1056, 330, W - M - 1056, 260, headline, size=64, weight=700, tracking=-1,
                   line=1.08, anchor="bottom")
        if context:
            self._text(s, 1056, 612, W - M - 1056, 90, context, size=28, color="slate", line=1.3)
        self._footer(s)
        return s

    def bullets(self, eyebrow, headline, bullets, *, aside=None, notes=None):
        """Takeaway headline + up to 3 short bullets (emerald dash). Optional
        aside = {"value", "label", "context"} as the slide's one dark highlight card."""
        s = self._slide("platinum", notes)
        self._eyebrow(s, eyebrow)
        self._headline(s, headline, w=1500 if not aside else 1100)
        text_w = 1050 if aside else 1500
        self._dash_list(s, bullets, M, CONTENT_TOP + 20, text_w)
        if aside:
            x, y, w, h = 1264, CONTENT_TOP, W - M - 1264, CONTENT_BOTTOM - CONTENT_TOP
            self._box(s, x, y, w, h, fill="prussian-blue", outline="prussian-blue")
            self._text(s, x + 48, y + 48, w - 96, 170, aside["value"], size=140, weight=900, tracking=-3,
                       color="emerald", line=1.0)
            self._text(s, x + 48, y + 236, w - 96, 30, aside.get("label", ""), size=20, weight=500,
                       caps=True, tracking=1, color="white")
            if aside.get("context"):
                self._text(s, x + 48, y + 284, w - 96, 200, aside["context"], size=26, color="white",
                           line=1.3)
        self._footer(s)
        return s

    def kpi_chart(self, eyebrow, headline, *, kpis, chart, note=None, number_format='#,##0',
                  notes=None):
        """Up to 4 KPI tiles left, one highlighted bar chart right, one annotation."""
        s = self._slide("platinum", notes)
        self._eyebrow(s, eyebrow)
        self._headline(s, headline)
        # KPI tiles
        kpis = kpis[:4]
        n = len(kpis)
        x0, y0, aw, g, th = M, CONTENT_TOP, 720, 24, 236
        cols = 1 if n <= 2 else 2
        rows = math.ceil(n / cols)
        for i, k in enumerate(kpis):
            r, c = divmod(i, cols)
            full = cols == 1 or (i == n - 1 and n % 2 == 1)
            tw = aw if full else (aw - g) / 2
            tx, ty = x0 + (0 if full else c * (tw + g)), y0 + r * (th + g)
            hi = bool(k.get("highlight"))
            fill = "prussian-blue" if hi else "white"
            ink = "white" if hi else "prussian-blue"
            vsize = min(72, (tw - 64) / (max(1, len(str(k["value"]))) * 0.64))
            self._box(s, tx, ty, tw, th, fill=fill)
            self._text(s, tx + 32, ty + 28, tw - 64, 84, str(k["value"]), size=vsize, weight=900,
                       tracking=-1.5, color="emerald" if hi else ink, line=1.0, anchor="bottom")
            self._text(s, tx + 32, ty + 124, tw - 64, 28, k["label"], size=18, weight=500, caps=True,
                       tracking=1, color=ink)
            if k.get("context"):
                self._text(s, tx + 32, ty + 160, tw - 64, th - 172, k["context"], size=24,
                           color="white" if hi else "slate", line=1.2)
        tiles_bottom = y0 + rows * th + (rows - 1) * g
        # Chart: its baseline lines up with the bottom of the KPI tiles
        note_lines = max(1, math.ceil(len(note or "") * 24 * 0.52 / 460))
        cx, cw = 864, W - M - 864
        cy = CONTENT_TOP + (note_lines * 30 + 40 if note else 0)
        base = max(tiles_bottom, cy + 330)
        ch = (base - cy) / 0.88
        cats, vals = chart["categories"], chart["values"]
        hi_i = chart.get("highlight", len(vals) - 1)
        cmp_i = chart.get("comparison")
        data = CategoryChartData()
        data.categories = cats
        data.add_series(chart.get("series_name", "Value"), vals)
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, E(cx), E(cy), E(cw), E(ch), data)
        c = gf.chart
        c.has_legend = False
        c.has_title = False
        c.font.name, c.font.size = font_for(500)[0], Pt(10)
        c.font.color.rgb = rgb("prussian-blue")
        plot = c.plots[0]
        plot.gap_width = 70
        plot.vary_by_categories = False
        ser = plot.series[0]
        for i in range(len(vals)):
            pt = ser.points[i]
            pt.format.fill.solid()
            tok = "emerald" if i == hi_i else ("full-spectrum-blue" if i == cmp_i else "blue-wash")
            pt.format.fill.fore_color.rgb = rgb(tok)
            pt.format.line.color.rgb = rgb("prussian-blue")
            pt.format.line.width = E(2)
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = number_format, False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size, dl.font.name = Pt(12), font_for(700)[0]
        dl.font.bold = True
        dl.font.color.rgb = rgb("prussian-blue")
        vmax = max(vals) * 1.2
        va = c.value_axis
        va.minimum_scale, va.maximum_scale = 0, vmax
        va.has_major_gridlines = va.has_minor_gridlines = False
        va.visible = False
        ca = c.category_axis
        ca.has_major_gridlines = False
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        ca.format.line.color.rgb = rgb("prussian-blue")
        ca.format.line.width = E(2)
        ca.tick_labels.font.size = Pt(10)
        ca.tick_labels.font.color.rgb = rgb("slate")
        # deterministic plot area so the annotation can point at the highlighted bar
        cs = c._chartSpace
        pa = cs.find(f".//{{{C_NS}}}plotArea")
        old = pa.find(f"{{{C_NS}}}layout")
        if old is not None:
            pa.remove(old)
        lay = etree.fromstring(
            f'<c:layout xmlns:c="{C_NS}"><c:manualLayout><c:layoutTarget val="inner"/>'
            '<c:xMode val="edge"/><c:yMode val="edge"/><c:x val="0.02"/><c:y val="0.04"/>'
            '<c:w val="0.96"/><c:h val="0.84"/></c:manualLayout></c:layout>')
        pa.insert(0, lay)
        sp = etree.fromstring(f'<c:spPr xmlns:c="{C_NS}" xmlns:a="{A_NS}"><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>')
        cs.find(f"{{{C_NS}}}chart").addnext(sp)
        if note:
            px, pw = cx + cw * 0.02, cw * 0.96
            py, ph = cy + ch * 0.04, ch * 0.84
            bx = px + (hi_i + 0.5) * pw / len(vals)
            btop = py + ph * (1 - vals[hi_i] / vmax) - 44
            nw = 460
            nx = bx - 40 - nw if bx + 40 + nw > W - M else bx + 40
            nx = max(nx, cx)
            ny = CONTENT_TOP
            nh = note_lines * 30
            self._text(s, nx, ny, nw, nh, note, size=24, color="prussian-blue", line=1.2,
                       align="right" if nx < bx else "left")
            ax = nx + nw if nx < bx else nx
            con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(ax), E(ny + nh / 2), E(bx),
                                         E(max(btop, ny + nh + 16)))
            _strip_style(con)
            con.line.color.rgb = rgb("prussian-blue")
            con.line.width = E(2)
        self._footer(s)
        return s

    def cards(self, eyebrow, headline, cards, *, notes=None):
        """2–5 cards in an asymmetric grid (some cards two columns wide).
        Each card: {"title", "body", "fill"?}. Default fills: white, emerald,
        white, full-spectrum-blue, white. Coral only as the one warm accent on a
        small card (set "fill": "vibrant-coral" yourself); max four fills per grid."""
        layouts = {
            2: [(0, 0, 2, 2), (2, 0, 1, 2)],
            3: [(0, 0, 2, 1), (2, 0, 1, 2), (0, 1, 2, 1)],
            4: [(0, 0, 2, 1), (2, 0, 1, 1), (0, 1, 1, 1), (1, 1, 2, 1)],
            5: [(0, 0, 2, 1), (2, 0, 1, 1), (0, 1, 1, 1), (1, 1, 1, 1), (2, 1, 1, 1)],
        }
        cards = cards[:5]
        if len(cards) < 2:
            raise ValueError("cards() needs 2–5 cards")
        fills = ["white", "emerald", "white", "full-spectrum-blue", "white"]
        s = self._slide("platinum", notes)
        self._eyebrow(s, eyebrow)
        self._headline(s, headline)
        gx, gy, gw, gh, g = M, CONTENT_TOP, W - 2 * M, CONTENT_BOTTOM - CONTENT_TOP, 24
        colw, rowh = (gw - 2 * g) / 3, (gh - g) / 2
        for idx, (card, (c0, r0, cs_, rs)) in enumerate(zip(cards, layouts[len(cards)])):
            x, y = gx + c0 * (colw + g), gy + r0 * (rowh + g)
            w, h = colw * cs_ + g * (cs_ - 1), rowh * rs + g * (rs - 1)
            fill = card.get("fill") or fills[idx % len(fills)]
            ink = text_on(fill)
            self._box(s, x, y, w, h, fill=fill)
            self._text(s, x + 40, y + 36, w - 80, 56, card["title"], size=40, weight=700, tracking=-0.6,
                       color=ink, line=1.1)
            if card.get("body"):
                self._text(s, x + 40, y + 104, w - 80, h - 130, card["body"], size=28, color=ink, line=1.3)
        self._footer(s)
        return s

    def story(self, eyebrow, headline, points, *, photo=None, notes=None):
        """Cutout photo on the left (duotone ground), headline + up to 3 takeaways right."""
        s = self._slide("platinum", notes)
        pw = 860
        if photo and Path(photo).exists():
            from PIL import Image
            iw, ih = Image.open(photo).size
            pic = s.shapes.add_picture(str(photo), 0, 0, E(pw), E(H))
            target, src = pw / H, iw / ih
            if src > target:
                cut = (1 - target / src) / 2
                pic.crop_left = pic.crop_right = cut
            else:
                cut = (1 - src / target) / 2
                pic.crop_top = pic.crop_bottom = cut
        else:
            self._box(s, 0, 0, pw, H, fill="full-spectrum-blue", outline=None, radius=0,
                      shape=MSO_SHAPE.RECTANGLE)
            self._text(s, 64, H / 2 - 60, pw - 128, 120,
                       "Photo: blue-duotone ground + cutout with white sticker outline (add in your design tool)",
                       size=24, color="white", align="center", anchor="middle", line=1.3)
        x = pw + 96
        self._eyebrow(s, eyebrow)
        s.shapes[-1].left = E(x)
        self._headline(s, headline, x=x, w=W - M - x, size=64)
        self._dash_list(s, points, x, CONTENT_TOP + 20, W - M - x, size=28)
        self._footer(s)
        # footer monogram must not sit on the photo: move it to the right column
        for shp in list(s.shapes)[-3:]:
            if shp.left < E(x):
                shp.left = shp.left + E(x - M)
        return s

    def closing(self, headline="Thank you", *, script=None, contact=None, notes=None):
        """Closing: prussian ground, 'Thank you' in Satoshi Black, script flourish,
        contact line, inline white logo."""
        s = self._slide("prussian-blue", notes)
        self._logo(s, "inline-white", M, M, w=400)
        self._blob(s, W - 430, H - 470, 660)
        self._text(s, M, 380, 1300, 190, headline, size=160, weight=900, tracking=-4, color="white",
                   line=1.0)
        if script:
            self._text(s, M + 12, 572, 900, 120, script, size=72, family="accent", color="emerald",
                       rotation=-3)
        if contact:
            self._text(s, M, 760, 1300, 50, contact, size=28, color="white")
        return s

    # ------------------------------------------------------------------ output
    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        if self._missing:
            print("WARNING: missing logo files, placeholders used:", ", ".join(sorted(self._missing)))
        print(f"Saved {path} ({len(self.prs.slides)} slides)")
        return path


# ---------------------------------------------------------------- previews
def install_fonts():
    """Make the brand fonts available to LibreOffice for previews (Linux/macOS)."""
    src = SKILL / "assets" / "fonts"
    dest = Path.home() / (".fonts" if sys.platform.startswith("linux") else "Library/Fonts")
    dest.mkdir(parents=True, exist_ok=True)
    added = False
    for f in src.glob("*.[ot]tf"):
        if not (dest / f.name).exists():
            shutil.copy(f, dest / f.name)
            added = True
    if added and shutil.which("fc-cache"):
        subprocess.run(["fc-cache", "-f"], capture_output=True)


def render_previews(pptx, out_dir, dpi=80):
    """PPTX -> PNG per slide via LibreOffice + pdftoppm, for visual checks."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not shutil.which("pdftoppm"):
        print("Preview skipped: LibreOffice/pdftoppm not available.")
        return []
    install_fonts()
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(pptx)],
                   check=True, capture_output=True, timeout=180)
    pdf = out / (Path(pptx).stem + ".pdf")
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(pdf), str(out / "slide")], check=True)
    pngs = sorted(out.glob("slide-*.png"))
    print(f"Previews: {len(pngs)} PNGs in {out}")
    return pngs


def build_from_spec(spec, out):
    d = Deck()
    for sl in spec["slides"]:
        sl = dict(sl)
        kind = sl.pop("type")
        getattr(d, kind)(**sl)
    return d.save(out)


DEMO = {"slides": [
    {"type": "title", "eyebrow": "Pitch deck", "headline": "Software built with",
     "headline_2": "brilliance and precision", "script": "for Acme", "period": "Q4 '26"},
    {"type": "section", "number": 1, "title": "The problem", "script": "why now"},
    {"type": "big_number", "eyebrow": "The cost of late QA", "number": "40%",
     "headline": "Releases slip by 40% when testing starts after build",
     "context": "[Source: add the study or client data behind this number]"},
    {"type": "bullets", "eyebrow": "Our approach", "headline": "We embed QA from the first sprint",
     "bullets": ["Test plans written alongside every user story",
                 "Automated checks run on every merge",
                 "A weekly quality report for your stakeholders"],
     "aside": {"value": "3.1x", "label": "Faster releases", "context": "[Client name], first six months"}},
    {"type": "section", "number": 2, "title": "Results", "script": "the proof"},
    {"type": "kpi_chart", "eyebrow": "Quality", "headline": "Defects fell 62% in six months",
     "kpis": [{"value": "62%", "label": "Fewer defects", "context": "▼ down vs. January", "highlight": True},
              {"value": "3.1x", "label": "Faster releases", "context": "▲ up, monthly to weekly"},
              {"value": "99.9%", "label": "Uptime", "context": "Across all environments"},
              {"value": "14", "label": "Engineers", "context": "Dedicated, one time zone"}],
     "chart": {"categories": ["Jan", "Feb", "Mar", "Apr", "May", "Jun"],
               "values": [120, 104, 88, 70, 58, 46], "highlight": 5},
     "note": "June: lowest defect count on record"},
    {"type": "cards", "eyebrow": "Services", "headline": "Four ways we help you ship",
     "cards": [{"title": "Product engineering", "body": "Web and mobile products, from first sprint to scale."},
               {"title": "QA and automation", "body": "Testing built into delivery."},
               {"title": "Cloud and DevOps", "body": "Pipelines and infrastructure that stay up."},
               {"title": "Dedicated teams", "body": "Engineers who join your stand-up from day one."}]},
    {"type": "section", "number": 3, "title": "Our team", "script": "the people"},
    {"type": "story", "eyebrow": "Culture", "headline": "Our engineers learn together every week",
     "points": ["Weekly soft-skills and tech sessions", "Mentors for every new joiner",
                "Teams that stay on your account"]},
    {"type": "closing", "script": "let's build together", "contact": "[contact email]  ·  codingcops.com"},
]}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", help="JSON file: {\"slides\": [{\"type\": ..., ...}]}")
    ap.add_argument("--out", help="Output .pptx")
    ap.add_argument("--demo", metavar="OUT", help="Build the demo deck to OUT")
    ap.add_argument("--preview", metavar="DIR", help="Also render PNG previews into DIR")
    a = ap.parse_args()
    if a.demo:
        out = build_from_spec(copy.deepcopy(DEMO), a.demo)
    elif a.spec and a.out:
        out = build_from_spec(json.loads(Path(a.spec).read_text(encoding="utf-8")), a.out)
    else:
        ap.error("use --spec SPEC --out OUT, or --demo OUT")
    if a.preview:
        render_previews(out, a.preview)


if __name__ == "__main__":
    main()
