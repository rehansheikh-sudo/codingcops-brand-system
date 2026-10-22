#!/usr/bin/env python3
"""
cc_docx.py — build on-brand CodingCops Word documents (A4), per guidelines/09-documents.md.

    import sys; sys.path.insert(0, "<skill>/scripts")
    from cc_docx import Doc

    # Report / proposal with a cover page
    d = Doc("PayLane payments app", doc_type="Proposal", classification="Confidential")
    d.cover(subtitle="Team and plan for v1", meta="October 2026 · v1.0 · Prepared by [Name]")
    d.h1("Summary")
    d.p("We propose a dedicated team of ...")
    d.key_figures([("40+", "Fintech products shipped"), ("2 weeks", "Team start"), ("Weekly", "Release cadence")])
    d.h2("Scope")
    d.bullets(["Payments app for iOS and Android", "Admin dashboard", "QA and automation"])
    d.table(["Phase", "Weeks", "Cost (USD)"], [["Discovery", "2", "[X]"], ["Build", "12", "[X]"]], numeric_cols=[1, 2])
    d.callout("Start in two weeks with a team that has shipped 40+ fintech products.", label="Key takeaway")
    d.save("CodingCops_Proposal_PayLane_2026-10.docx")

    # Letter or one-pager (letterhead layout, no cover)
    d = Doc("Proposal summary", kind="letter", address=["[Street address]", "[City]", "codingcops.com"])

    python3 cc_docx.py --demo out.docx --preview previews/
"""
from __future__ import annotations

import argparse
import copy
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor
from lxml import etree

SKILL = Path(__file__).resolve().parents[1]
TOKENS = json.loads((SKILL / "tokens" / "brand-tokens.json").read_text(encoding="utf-8"))
LOGOS = SKILL / TOKENS["logos"]["$dir"]
BLOB = json.loads((SKILL / "assets" / "shapes" / "emerald-blob.json").read_text(encoding="utf-8"))
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
PAGE_W, PAGE_H = 210, 297  # mm
TEXT_W = 170               # mm (20 mm margins)

FONT = {300: ("Satoshi Light", False), 400: ("Satoshi", False), 500: ("Satoshi Medium", False),
        700: ("Satoshi", True), 900: ("Satoshi Black", False)}


def hexof(tok):
    return TOKENS["colors"][tok]["hex"].lstrip("#").upper()


def rgb(tok):
    return RGBColor.from_string(hexof(tok))


def _set_font(rpr_owner_font, weight=400, size=None, color=None, caps=False, spacing=None):
    """Apply Satoshi weight/size/colour to a python-docx Font (run or style)."""
    name, bold = FONT[weight]
    f = rpr_owner_font
    f.name, f.bold = name, bold
    rpr = f.element.get_or_add_rPr() if hasattr(f.element, "get_or_add_rPr") else f._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn(f"w:{a}"), None)
    for a in ("ascii", "hAnsi", "cs"):
        rf.set(qn(f"w:{a}"), name)
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = rgb(color)
        c = rpr.find(qn("w:color"))
        for a in ("themeColor", "themeShade", "themeTint"):
            c.attrib.pop(qn(f"w:{a}"), None)
    f.all_caps = bool(caps)
    if spacing is not None:
        sp = rpr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            rpr.append(sp)
        sp.set(qn("w:val"), str(int(spacing * 20)))


def _border(el_pr, side, sz_eighths, color, space=0, tag="w:pBdr"):
    bdr = el_pr.find(qn(tag))
    if bdr is None:
        bdr = OxmlElement(tag)
        el_pr.append(bdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single" if sz_eighths else "nil")
    if sz_eighths:
        b.set(qn("w:sz"), str(sz_eighths))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), hexof(color))
    bdr.append(b)


def _field(paragraph, instr, placeholder, weight=400, size=8, color="slate"):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = paragraph.add_run(placeholder)
    _set_font(r.font, weight, size, color)
    r._r.getparent().remove(r._r)
    fld.append(r._r)
    paragraph._p.append(fld)


def _float_picture(paragraph, path, width_mm, height_mm, x_mm, y_mm, behind=True):
    """Insert a picture positioned on the page (not in the text flow)."""
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Mm(width_mm), height=Mm(height_mm))._inline
    anchor = OxmlElement("wp:anchor")
    for k, v in dict(distT="0", distB="0", distL="0", distR="0", simplePos="0", relativeHeight="0",
                     behindDoc="1" if behind else "0", locked="1", layoutInCell="1", allowOverlap="1").items():
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos"); sp.set("x", "0"); sp.set("y", "0"); anchor.append(sp)
    for tag, val in (("wp:positionH", x_mm), ("wp:positionV", y_mm)):
        pos = OxmlElement(tag); pos.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset"); off.text = str(int(Mm(val))); pos.append(off)
        anchor.append(pos)
    anchor.append(inline.find(qn("wp:extent")))
    ee = OxmlElement("wp:effectExtent")
    for k in "ltrb":
        ee.set(k, "0")
    anchor.append(ee)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(inline.find(qn("wp:docPr")))
    anchor.append(inline.find(qn("wp:cNvGraphicFramePr")))
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)
    return run


def _cover_background(path, dpi=200):
    """Full-bleed prussian A4 with the inline white logo (40 mm, top-left) and the
    emerald blob bleeding off the bottom-right corner."""
    from PIL import Image, ImageDraw
    mm = dpi / 25.4
    W, H = int(PAGE_W * mm), int(PAGE_H * mm)
    img = Image.new("RGB", (W, H), "#" + hexof("prussian-blue"))
    # blob, supersampled for smooth edges
    bh = 120 * mm
    bw = bh * BLOB["aspect_w_over_h"]
    bx, by = W - 0.62 * bw, H - 0.55 * bh
    ss = 4
    mask = Image.new("L", (int(bw * ss), int(bh * ss)), 0)
    ImageDraw.Draw(mask).polygon([(px * bw * ss, py * bh * ss) for px, py in BLOB["points_normalized"]], fill=255)
    mask = mask.resize((int(bw), int(bh)), Image.LANCZOS)
    img.paste(Image.new("RGB", mask.size, "#" + hexof("emerald")), (int(bx), int(by)), mask)
    logo = Image.open(LOGOS / TOKENS["logos"]["inline-white"]["file"]).convert("RGBA")
    lw = int(40 * mm)
    logo = logo.resize((lw, int(lw * logo.height / logo.width)), Image.LANCZOS)
    img.paste(logo, (int(20 * mm), int(20 * mm)), logo)
    img.save(path, optimize=True)
    return path


def _bar_png(path, token, w=2000, h=30):
    from PIL import Image
    Image.new("RGB", (w, h), "#" + hexof(token)).save(path)
    return path


class Doc:
    def __init__(self, title, *, doc_type="Report", classification="Confidential", kind="report",
                 address=None):
        """kind='report' (proposals, reports, policies; optional cover) or
        kind='letter' (letters and one-pagers: letterhead on page 1)."""
        self.title, self.doc_type, self.classification, self.kind = title, doc_type, classification, kind
        self.doc = Document()
        self._tmp = Path(tempfile.mkdtemp())
        self._has_cover = False
        self._started = False
        self._styles()
        self._theme()
        sec = self.doc.sections[0]
        self._page(sec)
        if kind == "letter":
            self._letterhead(sec, address or ["[Street address]", "[City, Country]", "codingcops.com"])
        else:
            self._interior_header_footer(sec)
        s = self.doc.settings.element
        if s.find(qn("w:embedTrueTypeFonts")) is None:
            s.insert(0, OxmlElement("w:embedTrueTypeFonts"))
        self.doc.core_properties.title = f"{doc_type}: {title}" if doc_type else title

    # ----------------------------------------------------------------- setup
    def _page(self, sec):
        sec.page_width, sec.page_height = Mm(PAGE_W), Mm(PAGE_H)
        sec.left_margin = sec.right_margin = sec.top_margin = Mm(20)
        sec.bottom_margin = Mm(25)
        sec.header_distance, sec.footer_distance = Mm(9), Mm(10)

    def _styles(self):
        st = self.doc.styles
        keep = {"Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption", "Header", "Footer",
                "Default Paragraph Font", "Normal Table", "No List", "Table Grid",
                "Title Char", "Heading 1 Char", "Heading 2 Char", "Heading 3 Char", "Header Char", "Footer Char"}
        for s in list(st):
            if s.name not in keep:
                s.element.getparent().remove(s.element)
        dd = st.element.find(qn("w:docDefaults"))
        rf = dd.find(".//" + qn("w:rFonts"))
        for a in list(rf.attrib):
            rf.attrib.pop(a)
        for a in ("ascii", "hAnsi", "cs"):
            rf.set(qn(f"w:{a}"), "Satoshi")

        def para(name, weight, size, color="prussian-blue", before=0, after=6, line=1.5, caps=False,
                 spacing=None, keep_next=False):
            s = st[name] if name in [x.name for x in st] else st.add_style(name, 1)
            _set_font(s.font, weight, size, color, caps, spacing)
            ppr = s.element.get_or_add_pPr()
            for tag in ("w:pBdr", "w:shd"):  # drop Word's default Office-blue rules/fills
                old = ppr.find(qn(tag))
                if old is not None:
                    ppr.remove(old)
            pf = s.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.line_spacing = line
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.keep_with_next = keep_next
            return s

        para("Normal", 400, 10.5)
        para("Title", 900, 40, line=1.0, after=8)
        h1 = para("Heading 1", 700, 24, before=18, after=8, line=1.1, keep_next=True)
        _border(h1.element.get_or_add_pPr(), "top", 24, "emerald", space=12)
        para("Heading 2", 700, 15, before=14, after=4, line=1.15, keep_next=True)
        para("Heading 3", 700, 11.5, before=10, after=2, line=1.2, keep_next=True)
        para("Caption", 400, 8.5, "slate", before=4, after=10, line=1.3)
        para("Eyebrow", 500, 8, "full-spectrum-blue", after=4, line=1.2, caps=True, spacing=0.5, keep_next=True)
        for name in ("Header", "Footer"):
            s_ = para(name, 500 if name == "Header" else 400, 8, "slate", after=0, line=1.0)
            tabs = s_.element.get_or_add_pPr().find(qn("w:tabs"))
            if tabs is not None:  # drop Word's default centre/right tabs
                tabs.getparent().remove(tabs)
        for name in ("Title Char", "Heading 1 Char", "Heading 2 Char", "Heading 3 Char"):
            if name in [x.name for x in st]:
                _set_font(st[name].font, {"Title Char": 900}.get(name, 700), None, "prussian-blue")
        for name in ("Header Char", "Footer Char"):
            if name in [x.name for x in st]:
                _set_font(st[name].font, 400, None, "slate")

    def _theme(self):
        for rel in self.doc.part.rels.values():
            if rel.reltype == RT.THEME:
                part = rel.target_part
                root = etree.fromstring(part.blob)
                ns = {"a": A_NS}
                for tag in ("majorFont", "minorFont"):
                    root.find(f".//a:{tag}/a:latin", ns).set("typeface", "Satoshi")
                scheme = {"dk1": "prussian-blue", "lt1": "white", "dk2": "prussian-blue", "lt2": "platinum",
                          "accent1": "emerald", "accent2": "full-spectrum-blue", "accent3": "vibrant-coral",
                          "accent4": "slate", "accent5": "emerald-wash", "accent6": "blue-wash",
                          "hlink": "full-spectrum-blue", "folHlink": "full-spectrum-blue"}
                clr = root.find(".//a:clrScheme", ns)
                for t, tok in scheme.items():
                    el = clr.find(f"a:{t}", ns)
                    for ch in list(el):
                        el.remove(ch)
                    etree.SubElement(el, f"{{{A_NS}}}srgbClr", val=hexof(tok))
                part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _interior_header_footer(self, sec):
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.style = self.doc.styles["Header"]
        hp.text = ""
        mono = LOGOS / TOKENS["logos"]["monogram"]["file"]
        hp.add_run().add_picture(str(mono), height=Mm(6))
        hp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        r = hp.add_run("\u2003" + self.title)
        _set_font(r.font, 500, 8, "full-spectrum-blue", caps=True, spacing=0.5)
        r = hp.add_run("\t")
        _field(hp, 'STYLEREF "Heading 1"', "Section", 500, 8, "slate")
        _border(hp._p.get_or_add_pPr(), "bottom", 4, "mist-border", space=6)
        fp = sec.footer.paragraphs[0]
        fp.style = self.doc.styles["Footer"]
        fp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        r = fp.add_run(f"CodingCops · {self.classification}\tPage ")
        _set_font(r.font, 400, 8, "slate")
        _field(fp, "PAGE", "1")
        r = fp.add_run(" of ")
        _set_font(r.font, 400, 8, "slate")
        _field(fp, "NUMPAGES", "1")

    def _letterhead(self, sec, address):
        """Page 1: inline blue logo (35 mm) left, address right (8 pt slate), 3 pt
        emerald rule at the bottom edge. Page 2+: monogram only."""
        sec.different_first_page_header_footer = True
        fh = sec.first_page_header
        t = fh.add_table(1, 2, Mm(TEXT_W))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        c0, c1 = t.rows[0].cells
        p0 = c0.paragraphs[0]
        p0.add_run().add_picture(str(LOGOS / TOKENS["logos"]["inline-blue"]["file"]), width=Mm(35))
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for i, line in enumerate(address):
            r = p1.add_run(line)
            _set_font(r.font, 400, 8, "slate")
            if i < len(address) - 1:
                r.add_break()
        fh.paragraphs[0]._p.getparent().remove(fh.paragraphs[0]._p) if len(fh.paragraphs) > 1 else None
        ff = sec.first_page_footer
        bar = _bar_png(self._tmp / "emerald-bar.png", "emerald")
        _float_picture(ff.paragraphs[0], bar, PAGE_W, 1.06, 0, PAGE_H - 1.06)
        hp = sec.header.paragraphs[0]
        hp.add_run().add_picture(str(LOGOS / TOKENS["logos"]["monogram"]["file"]), height=Mm(6))

    # ----------------------------------------------------------------- cover
    def cover(self, *, subtitle=None, meta=None, script=None):
        """Full-bleed prussian cover (reports and proposals longer than ~2 pages).
        Call first, before any content."""
        if self.kind != "report" or self._started:
            raise RuntimeError("cover() is for kind='report' and must be called before any content")
        self._has_cover = True
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True  # no header/footer on the cover
        first = self.doc.add_paragraph()
        first.paragraph_format.space_after = Pt(0)
        bg = _cover_background(self._tmp / "cover.png")
        _float_picture(first, bg, PAGE_W, PAGE_H, 0, 0)
        right = Mm(62)  # keep text clear of the blob

        def line(text, weight, size, color, after=4, caps=False, spacing=None, family=None, ls=1.1):
            p = self.doc.add_paragraph()
            # prussian paragraph shading: invisible on the cover image, keeps text
            # readable if the image is removed, and tells checkers the ground colour
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), hexof("prussian-blue")); ppr.append(shd)
            p.paragraph_format.right_indent = right
            p.paragraph_format.space_after = Pt(after)
            p.paragraph_format.line_spacing = ls
            r = p.add_run(text)
            if family == "accent":
                r.font.name = "Brittany Signature"
                rf = r._r.get_or_add_rPr().find(qn("w:rFonts"))
                for a in ("ascii", "hAnsi", "cs"):
                    rf.set(qn(f"w:{a}"), "Brittany Signature")
                r.font.size, r.font.color.rgb = Pt(size), rgb(color)
            else:
                _set_font(r.font, weight, size, color, caps, spacing)
            return p

        # Push the title block to the bottom-left: the anchor paragraph becomes an
        # exact-height spacer (works in Word, LibreOffice and Google Docs alike).
        import math
        title_lines = max(1, math.ceil(len(self.title) / 11))
        # rendered line heights: Satoshi sets ~1.35x its size even at "single" spacing
        block = 19 + title_lines * 60 + 6 + (51 if script else 0) + (33 if subtitle else 0) + 14
        usable = (PAGE_H - 20 - 25) * 72 / 25.4
        first.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        first.paragraph_format.line_spacing = Pt(max(12, usable - block - 30))
        line(self.doc_type, 500, 8, "emerald", after=8, caps=True, spacing=0.5)
        line(self.title, 900, 44, "white", after=6, ls=1.0)
        if script:
            line(script, 400, 28, "white", after=6, family="accent")
        if subtitle:
            line(subtitle, 500, 14, "emerald", after=14)
        line(meta or f"[Month YYYY] · v1.0 · Prepared by [Name] · {self.classification}", 400, 8.5, "white", after=0)
        new = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(new)
        new.different_first_page_header_footer = False
        nva = new._sectPr.find(qn("w:vAlign"))
        if nva is not None:
            new._sectPr.remove(nva)
        self._interior_header_footer(new)
        return self

    # --------------------------------------------------------------- content
    def _para(self, text="", style="Normal"):
        self._started = True
        return self.doc.add_paragraph(text, style=style)

    def title_block(self, title=None, eyebrow=None, subtitle=None):
        """Title for documents without a cover (one-pagers, letters, short reports)."""
        if eyebrow:
            self._para(eyebrow, "Eyebrow")
        self._para(title or self.title, "Title").paragraph_format.space_after = Pt(4)
        if subtitle:
            p = self._para("", "Normal")
            r = p.add_run(subtitle)
            _set_font(r.font, 500, 13, "slate")
        return self

    def eyebrow(self, text):
        self._para(text, "Eyebrow"); return self

    def h1(self, text, new_page=False):
        p = self._para(text, "Heading 1")
        p.paragraph_format.page_break_before = new_page
        return self

    def h2(self, text):
        self._para(text, "Heading 2"); return self

    def h3(self, text):
        self._para(text, "Heading 3"); return self

    def p(self, text):
        self._para(text, "Normal"); return self

    def caption(self, text):
        self._para(text, "Caption"); return self

    def bullets(self, items):
        """Emerald-dash list (max 7 items; use a table beyond that)."""
        for it in items[:7]:
            p = self._para("", "Normal")
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Mm(5), Mm(-5)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Mm(5))
            r = p.add_run("–\t")
            _set_font(r.font, 700, 10.5, "emerald")
            r = p.add_run(it)
            _set_font(r.font, 400, 10.5, "prussian-blue")
        return self

    def _cell_fmt(self, cell, fill=None, borders=None, margins_mm=(2, 2.5)):
        tcPr = cell._tc.get_or_add_tcPr()
        if fill:
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), hexof(fill)); tcPr.append(shd)
        if borders:
            for side, (sz, col) in borders.items():
                _border(tcPr, side, sz, col, tag="w:tcBorders")
        mar = OxmlElement("w:tcMar")
        for side, v in (("top", margins_mm[0]), ("bottom", margins_mm[0]), ("left", margins_mm[1]), ("right", margins_mm[1])):
            m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), str(int(v * 56.7))); m.set(qn("w:type"), "dxa"); mar.append(m)
        tcPr.append(mar)

    def _no_table_borders(self, table):
        tblPr = table._tbl.tblPr
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            _border(tblPr, side, 0, "white", tag="w:tblBorders")

    def table(self, header, rows, *, numeric_cols=(), widths_mm=None, caption=None):
        """Header row bold 9 pt on blue-wash with a 1 pt prussian rule; 0.5 pt
        mist-border row rules; no vertical lines; numbers right-aligned."""
        self._started = True
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        self._no_table_borders(t)
        for ri, data in enumerate([header] + rows):
            for ci, val in enumerate(data):
                cell = t.rows[ri].cells[ci]
                if ri == 0:
                    self._cell_fmt(cell, "blue-wash", {"bottom": (8, "prussian-blue")})
                else:
                    self._cell_fmt(cell, None, {"bottom": (4, "mist-border")})
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                if ci in numeric_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(str(val))
                _set_font(r.font, 700 if ri == 0 else 400, 9 if ri == 0 else 9.5, "prussian-blue")
        if widths_mm:
            for row in t.rows:
                for cell, w in zip(row.cells, widths_mm):
                    cell.width = Mm(w)
        if caption:
            self.caption(caption)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return self

    def callout(self, text, *, label="Key takeaway", dark=False):
        """Boxed note with an eyebrow label. dark=True (prussian fill) once per
        spread for the main takeaway. Word can't round table corners; PDFs made
        from HTML should use the 4 mm radius."""
        self._started = True
        t = self.doc.add_table(rows=1, cols=1)
        self._no_table_borders(t)
        cell = t.rows[0].cells[0]
        line = (8, "prussian-blue")
        self._cell_fmt(cell, "prussian-blue" if dark else "white",
                       {s: line for s in ("top", "left", "bottom", "right")}, margins_mm=(4, 5))
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label)
        _set_font(r.font, 500, 8, "emerald" if dark else "full-spectrum-blue", caps=True, spacing=0.5)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(text)
        _set_font(r.font, 500, 11.5, "white" if dark else "prussian-blue")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return self

    def key_figures(self, figures, *, highlight=0):
        """Up to 3 figures in a row: Satoshi Black 28 pt + eyebrow label. The
        highlighted one gets the page's emerald accent (a 3 pt emerald rule above;
        never emerald text on white)."""
        self._started = True
        figures = figures[:3]
        t = self.doc.add_table(rows=1, cols=len(figures))
        self._no_table_borders(t)
        for i, (val, lab) in enumerate(figures):
            cell = t.rows[0].cells[i]
            top = (24, "emerald") if i == highlight else (4, "mist-border")
            self._cell_fmt(cell, None, {"top": top}, margins_mm=(3, 0.5))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(val))
            _set_font(r.font, 900, 28, "prussian-blue")
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            r = p2.add_run(lab)
            _set_font(r.font, 500, 8, "full-spectrum-blue", caps=True, spacing=0.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return self

    def page_break(self):
        self._para().add_run().add_break(WD_BREAK.PAGE); return self

    def signoff(self, name, role=None, closing="Kind regards,"):
        self.p(closing)
        p = self._para("", "Normal")
        p.paragraph_format.space_before = Pt(18)
        r = p.add_run(name); _set_font(r.font, 700, 10.5, "prussian-blue")
        if role:
            r = p.add_run("\n" + role); _set_font(r.font, 400, 9.5, "slate")
        return self

    # ---------------------------------------------------------------- output
    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        print(f"Saved {path}")
        return path


def to_pdf(docx_path, out_dir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("PDF skipped: LibreOffice not available.")
        return None
    sys.path.insert(0, str(Path(__file__).parent))
    from cc_deck import install_fonts
    install_fonts()
    out = Path(out_dir); out.mkdir(parents=True, exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx_path)],
                   check=True, capture_output=True, timeout=180)
    return out / (Path(docx_path).stem + ".pdf")


def render_previews(docx_path, out_dir, dpi=70):
    pdf = to_pdf(docx_path, out_dir)
    if not pdf or not shutil.which("pdftoppm"):
        return []
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(pdf), str(Path(out_dir) / "page")], check=True)
    pngs = sorted(Path(out_dir).glob("page-*.png"))
    print(f"Previews: {len(pngs)} PNGs in {out_dir}")
    return pngs


def demo(out):
    d = Doc("PayLane payments app", doc_type="Proposal")
    d.cover(subtitle="A dedicated team to build and launch v1",
            meta="October 2026 · v1.0 · Prepared by [Name] · Confidential")
    d.h1("Summary")
    d.p("You need a payments app that launches safely and keeps shipping. We have built fintech "
        "products for more than 40 clients, and our dedicated teams start within two weeks.")
    d.key_figures([("40+", "Fintech products shipped"), ("2 weeks", "To a working team"),
                   ("Weekly", "Release cadence")])
    d.h2("What we will build")
    d.bullets(["Customer payments app for iOS and Android", "Admin and support dashboard",
               "Automated testing from the first sprint"])
    d.h2("Plan and budget")
    d.table(["Phase", "Weeks", "Team", "Cost (USD)"],
            [["Discovery", "2", "3", "[X]"], ["Build v1", "12", "6", "[X]"], ["Launch and hand-over", "2", "4", "[X]"]],
            numeric_cols=[1, 2, 3], caption="Figure 1. Indicative plan. Source: CodingCops estimate, [date].")
    d.callout("A previous fintech client moved from monthly to weekly releases with this model.",
              label="Key takeaway", dark=True)
    d.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--demo", metavar="OUT", help="Build a demo proposal to OUT")
    ap.add_argument("--preview", metavar="DIR", help="Render PDF + PNG previews into DIR")
    a = ap.parse_args()
    if not a.demo:
        ap.error("use --demo OUT [--preview DIR], or import Doc from Python")
    out = demo(a.demo)
    if a.preview:
        render_previews(out, a.preview)


if __name__ == "__main__":
    main()
